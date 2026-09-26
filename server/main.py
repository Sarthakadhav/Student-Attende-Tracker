"""
CSE Duty Leave & Attendance Portal - Python backend (FastAPI).

Implements the Registrar's circular:
  * students need a pre-approval letter and a certificate (or HoD/Dean permission)
  * medical cases are recorded but not counted
  * Final % = (ERP attended + lectures missed during the event) / ERP total

Run:  python main.py      (http://localhost:5000, docs at http://localhost:5000/docs)
"""

import csv
import io
import json
import logging
import os
import re
import secrets
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

import bcrypt
import jwt
import uvicorn
from fastapi import Body, Depends, FastAPI, File, Form, Header, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from starlette.exceptions import HTTPException

import erp_matrix
import lecture_engine as engine
from database import UPLOAD_DIR, get_db, init_db

# =========================================
# CONFIG
# =========================================

PORT = int(os.getenv("PORT", "5000"))
CLIENT_URL = os.getenv("CLIENT_URL", "http://localhost:5173")
JWT_SECRET = os.getenv("JWT_SECRET", "dev-only-secret-change-me")
FACULTY_SIGNUP_CODE = os.getenv("FACULTY_SIGNUP_CODE", "CSE-FACULTY-2026")
TOKEN_DAYS = 7
THRESHOLD = 75.0

EVENT_CATEGORIES = [
    "Hackathon / Competition",
    "Technical event / Workshop",
    "Sports",
    "Cultural",
    "Internship",
    "NSS / NCC",
    "Other college duty",
]
EVIDENCE_TYPES = {
    "certificate": "Participation certificate with the exact dates",
    "authority_permission": "Permission signed by HoD / Dean (no certificate)",
}
ALLOWED_FILE_TYPES = {"application/pdf": ".pdf", "image/jpeg": ".jpg", "image/png": ".png"}
MAX_FILE_SIZE = 5 * 1024 * 1024
EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
PASSWORD_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghjkmnpqrstuvwxyz23456789"

log = logging.getLogger("attendance")
logging.basicConfig(level=logging.INFO)

if "JWT_SECRET" not in os.environ:
    log.warning("JWT_SECRET is not set. Using a development secret. Set it before deploying.")

# =========================================
# APP
# =========================================

app = FastAPI(title="CSE Duty Leave & Attendance API")
app.add_middleware(
    CORSMiddleware, allow_origins=[CLIENT_URL], allow_methods=["*"], allow_headers=["*"]
)
init_db()


class ApiError(HTTPException):
    """An error whose message is safe to show to the user."""


@app.exception_handler(HTTPException)
async def http_error(_request: Request, exc: HTTPException):
    message = exc.detail
    if exc.status_code == 404 and message == "Not Found":
        message = "API route not found."
    elif exc.status_code == 405:
        message = "This method isn't allowed on this route."
    return JSONResponse({"error": message}, status_code=exc.status_code)


@app.exception_handler(RequestValidationError)
async def validation_error(_request: Request, exc: RequestValidationError):
    json_problem = any(e.get("type") == "json_invalid" for e in exc.errors())
    message = "Request body is not valid JSON." if json_problem else "Some of the data sent is invalid."
    return JSONResponse({"error": message}, status_code=400)


@app.exception_handler(engine.EngineError)
async def engine_error(_request: Request, exc: engine.EngineError):
    return JSONResponse({"error": str(exc)}, status_code=400)


@app.exception_handler(Exception)
async def server_error(_request: Request, exc: Exception):
    log.exception("Unhandled error", exc_info=exc)
    return JSONResponse({"error": "Something went wrong on the server."}, status_code=500)


# =========================================
# HELPERS
# =========================================


def clean(value) -> str:
    return value.strip() if isinstance(value, str) else ""


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt(rounds=10)).decode()


def check_password(password: str, password_hash: str) -> bool:
    if len(password.encode("utf-8")) > 72:
        return False
    return bcrypt.checkpw(password.encode("utf-8"), password_hash.encode("utf-8"))


def generate_password() -> str:
    return "".join(secrets.choice(PASSWORD_ALPHABET) for _ in range(10))


def validate_new_password(password: str):
    if len(password) < 8:
        raise ApiError(400, "Password must be at least 8 characters.")
    if len(password.encode("utf-8")) > 72:
        raise ApiError(400, "Password is too long (maximum 72 characters).")


def faculty_class_ids(conn, faculty_id: str) -> list[str]:
    rows = conn.execute("SELECT class_id FROM faculty_classes WHERE faculty_id = ?", (faculty_id,))
    return [r["class_id"] for r in rows]


def class_label(conn, class_id: str | None) -> str | None:
    if not class_id:
        return None
    row = conn.execute("SELECT label FROM classes WHERE id = ?", (class_id,)).fetchone()
    return row["label"] if row else class_id


def public_user(row, conn) -> dict:
    user = {
        "id": row["id"],
        "role": row["role"],
        "name": row["name"],
        "email": row["email"],
        "createdAt": row["created_at"],
        "prn": row["prn"],
        "rollNo": row["roll_no"],
        "classId": row["class_id"],
        "classLabel": class_label(conn, row["class_id"]),
        "designation": row["designation"],
        "mustChangePassword": bool(row["must_change_password"]),
    }
    if row["role"] == "faculty":
        user["classes"] = faculty_class_ids(conn, row["id"])
    return {k: v for k, v in user.items() if v is not None}


def sign_token(user_id: str) -> str:
    payload = {"sub": user_id, "exp": datetime.now(timezone.utc) + timedelta(days=TOKEN_DAYS)}
    return jwt.encode(payload, JWT_SECRET, algorithm="HS256")


def logged_in_user(authorization: str | None = Header(default=None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise ApiError(401, "Please log in to continue.")
    try:
        payload = jwt.decode(authorization[7:], JWT_SECRET, algorithms=["HS256"])
    except jwt.PyJWTError:
        raise ApiError(401, "Your session has expired. Please log in again.")

    with get_db() as conn:
        row = conn.execute("SELECT * FROM users WHERE id = ?", (payload.get("sub"),)).fetchone()
    if row is None:
        raise ApiError(401, "This account no longer exists.")
    return row


def current_user(user=Depends(logged_in_user)):
    """Like logged_in_user, but blocks everything until a starting password is changed."""
    if user["must_change_password"]:
        raise ApiError(403, "Change your starting password first.")
    return user


def require_role(role: str):
    def check(user=Depends(current_user)):
        if user["role"] != role:
            raise ApiError(403, "You don't have access to this.")
        return user

    return check


def faculty_can_see(conn, faculty, class_id: str) -> bool:
    allowed = faculty_class_ids(conn, faculty["id"])
    return not allowed or class_id in allowed  # no assigned classes = HoD, sees every class


def visible_classes(conn, faculty) -> list[str]:
    allowed = faculty_class_ids(conn, faculty["id"])
    if allowed:
        return allowed
    return [r["id"] for r in conn.execute("SELECT id FROM classes ORDER BY rowid")]


def student_info(row) -> dict:
    if row is None:
        return {"name": "Deleted student", "prn": "-", "classId": "-"}
    return {"name": row["name"], "prn": row["prn"], "rollNo": row["roll_no"], "classId": row["class_id"]}


def documents_for(conn, application_id: str) -> list[dict]:
    rows = conn.execute(
        "SELECT * FROM documents WHERE application_id = ? ORDER BY kind", (application_id,)
    ).fetchall()
    return [
        {"id": r["id"], "kind": r["kind"], "originalName": r["original_name"],
         "mimeType": r["mime_type"], "size": r["size"]}
        for r in rows
    ]


def application_json(row, conn) -> dict:
    student = conn.execute("SELECT * FROM users WHERE id = ?", (row["student_id"],)).fetchone()
    lectures = json.loads(row["lectures"])
    data = {
        "id": row["id"],
        "studentId": row["student_id"],
        "student": student_info(student),
        "classId": row["class_id"],
        "kind": row["kind"],
        "category": row["category"],
        "eventName": row["event_name"],
        "description": row["description"],
        "startDate": row["start_date"],
        "endDate": row["end_date"],
        "lectures": lectures,
        "sessions": sum(l["sessions"] for l in lectures),
        "evidenceType": row["evidence_type"],
        "certificatePending": bool(row["certificate_pending"]),
        "studentUnread": bool(row["student_unread"]),
        "reminderSentAt": row["reminder_sent_at"],
        "status": row["status"],
        "unread": bool(row["unread"]),
        "remark": row["remark"],
        "submittedAt": row["submitted_at"],
        "reviewedAt": row["reviewed_at"],
        "documents": documents_for(conn, row["id"]),
    }
    if row["reviewed_by"]:
        data["reviewedBy"] = row["reviewed_by"]
    return data


def get_application(conn, application_id: str):
    row = conn.execute("SELECT * FROM applications WHERE id = ?", (application_id,)).fetchone()
    if row is None:
        raise ApiError(404, "Application not found.")
    return row


def check_application_access(conn, user, row):
    if user["role"] == "student":
        if row["student_id"] != user["id"]:
            raise ApiError(403, "You don't have access to this application.")
    elif not faculty_can_see(conn, user, row["class_id"]):
        raise ApiError(403, "This application belongs to a class you don't coordinate.")


def subject_summary(class_id: str, lectures: list[dict]) -> list[dict]:
    """Sessions per subject for one application, split by CIA phase."""
    phases = engine.phases_for(class_id)
    rows = {}
    for lecture in lectures:
        cid = lecture["courseId"]
        row = rows.setdefault(cid, {
            "courseId": cid, "subjectCode": lecture["subjectCode"],
            "subjectName": lecture["subjectName"], "sessions": 0,
            "byPhase": {p["key"]: 0 for p in phases},
        })
        row["sessions"] += lecture["sessions"]
        key = engine.phase_key(phases, lecture["date"])
        if key:
            row["byPhase"][key] += lecture["sessions"]
    order = [c["id"] for c in engine.courses_for(class_id)]
    return sorted(rows.values(), key=lambda r: order.index(r["courseId"]) if r["courseId"] in order else 99)


def approved_event_lectures(conn, class_id: str) -> dict[str, list[dict]]:
    """student_id -> lectures from approved EVENT applications (medical never counts)."""
    result: dict[str, list[dict]] = {}
    rows = conn.execute(
        """SELECT student_id, lectures FROM applications
           WHERE class_id = ? AND kind = 'event' AND status = 'approved'""",
        (class_id,),
    ).fetchall()
    for r in rows:
        result.setdefault(r["student_id"], []).extend(json.loads(r["lectures"]))
    return result


def pending_counts(conn, class_id: str) -> dict[str, int]:
    rows = conn.execute(
        """SELECT student_id, COUNT(*) AS n FROM applications
           WHERE class_id = ? AND kind = 'event' AND status = 'pending' GROUP BY student_id""",
        (class_id,),
    ).fetchall()
    return {r["student_id"]: r["n"] for r in rows}


def class_students(conn, class_id: str):
    return conn.execute(
        "SELECT * FROM users WHERE role = 'student' AND class_id = ? ORDER BY roll_no, name",
        (class_id,),
    ).fetchall()


def require_class(conn, user, class_id: str):
    if not conn.execute("SELECT 1 FROM classes WHERE id = ?", (class_id,)).fetchone():
        raise ApiError(404, f"Unknown class {class_id}.")
    if not faculty_can_see(conn, user, class_id):
        raise ApiError(403, "You don't coordinate this class.")


def pct(numerator: float, denominator: float) -> float | None:
    return round(numerator / denominator * 100, 2) if denominator else None


async def read_upload(upload: UploadFile | None, label: str):
    if upload is None or not upload.filename:
        raise ApiError(400, f"Upload the {label}.")
    if upload.content_type not in ALLOWED_FILE_TYPES:
        raise ApiError(400, f"The {label} must be a PDF, JPG or PNG file.")
    content = await upload.read(MAX_FILE_SIZE + 1)
    if len(content) > MAX_FILE_SIZE:
        raise ApiError(400, f"The {label} must be smaller than 5 MB.")
    if not content:
        raise ApiError(400, f"The {label} file is empty.")
    return content, upload


def validate_dates(class_id: str, start_date, end_date) -> dict:
    for value in (start_date, end_date):
        if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ApiError(400, "Choose a valid start and end date.")
    try:
        return engine.missed_lectures(class_id, start_date, end_date)
    except ValueError:
        raise ApiError(400, "Choose a valid start and end date.")


# =========================================
# AUTH
# =========================================


@app.get("/api/classes")
def list_classes():
    """Public: used by the faculty sign-up form."""
    with get_db() as conn:
        rows = conn.execute("SELECT id, label FROM classes ORDER BY rowid").fetchall()
    return {"classes": [{"id": r["id"], "label": r["label"], "ready": engine.is_ready(r["id"])}
                        for r in rows]}


@app.post("/api/auth/login")
def login(body: dict = Body(default_factory=dict)):
    identifier = clean(body.get("identifier")).lower()
    password = body.get("password") if isinstance(body.get("password"), str) else ""
    if not identifier or not password:
        raise ApiError(400, "Enter your PRN or email and your password.")

    with get_db() as conn:
        row = conn.execute(
            "SELECT * FROM users WHERE LOWER(email) = ? OR LOWER(prn) = ?", (identifier, identifier)
        ).fetchone()
        # Same message either way, so nobody can find out which accounts exist.
        if row is None or not check_password(password, row["password_hash"]):
            raise ApiError(401, "Incorrect PRN/email or password.")
        return {"token": sign_token(row["id"]), "user": public_user(row, conn)}


@app.post("/api/auth/signup", status_code=201)
def faculty_signup(body: dict = Body(default_factory=dict)):
    """Only faculty sign up. Students are loaded from the department list (import_data.py)."""
    name = clean(body.get("name"))
    email = clean(body.get("email")).lower()
    password = body.get("password") if isinstance(body.get("password"), str) else ""
    designation = clean(body.get("designation")) or "Faculty"
    class_ids = body.get("classes") if isinstance(body.get("classes"), list) else []

    if clean(body.get("facultyCode")) != FACULTY_SIGNUP_CODE:
        raise ApiError(403, "The faculty access code is incorrect. Ask the department admin for it.")
    if len(name) < 2:
        raise ApiError(400, "Enter your full name.")
    if not EMAIL_RE.match(email):
        raise ApiError(400, "Enter a valid email address.")
    validate_new_password(password)

    with get_db() as conn:
        if conn.execute("SELECT 1 FROM users WHERE LOWER(email) = ?", (email,)).fetchone():
            raise ApiError(409, "An account with this email already exists.")
        known = {r["id"] for r in conn.execute("SELECT id FROM classes")}
        if any(c not in known for c in class_ids):
            raise ApiError(400, "One of the selected classes doesn't exist.")

        user_id = str(uuid.uuid4())
        conn.execute(
            """INSERT INTO users (id, role, name, email, password_hash, designation, created_at)
               VALUES (?, 'faculty', ?, ?, ?, ?, ?)""",
            (user_id, name, email, hash_password(password), designation, now_iso()),
        )
        conn.executemany(
            "INSERT INTO faculty_classes (faculty_id, class_id) VALUES (?, ?)",
            [(user_id, c) for c in dict.fromkeys(class_ids)],
        )
        row = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        return {"token": sign_token(user_id), "user": public_user(row, conn)}


@app.get("/api/auth/me")
def me(user=Depends(logged_in_user)):
    with get_db() as conn:
        return {"user": public_user(user, conn)}


@app.post("/api/auth/change-password")
def change_password(body: dict = Body(default_factory=dict), user=Depends(logged_in_user)):
    current = body.get("currentPassword") if isinstance(body.get("currentPassword"), str) else ""
    new = body.get("newPassword") if isinstance(body.get("newPassword"), str) else ""

    if not check_password(current, user["password_hash"]):
        raise ApiError(400, "Your current password is incorrect.")
    validate_new_password(new)
    if new == current:
        raise ApiError(400, "Choose a password different from the current one.")

    with get_db() as conn:
        conn.execute(
            "UPDATE users SET password_hash = ?, must_change_password = 0 WHERE id = ?",
            (hash_password(new), user["id"]),
        )
        row = conn.execute("SELECT * FROM users WHERE id = ?", (user["id"],)).fetchone()
        return {"user": public_user(row, conn)}


@app.post("/api/students/{prn}/reset-password")
def reset_student_password(prn: str, faculty=Depends(require_role("faculty"))):
    """Coordinator resets a forgotten password. The student must change it on next login."""
    with get_db() as conn:
        student = conn.execute(
            "SELECT * FROM users WHERE role = 'student' AND prn = ?", (prn.upper(),)
        ).fetchone()
        if student is None:
            raise ApiError(404, "No student with this PRN.")
        if not faculty_can_see(conn, faculty, student["class_id"]):
            raise ApiError(403, "This student is in a class you don't coordinate.")

        password = generate_password()
        conn.execute(
            "UPDATE users SET password_hash = ?, must_change_password = 1 WHERE id = ?",
            (hash_password(password), student["id"]),
        )
    return {"prn": student["prn"], "name": student["name"], "password": password}


# =========================================
# ACADEMIC DATA + ERP ATTENDANCE
# =========================================


@app.get("/api/academic")
def academic(user=Depends(current_user)):
    data = {"categories": EVENT_CATEGORIES, "evidenceTypes": EVIDENCE_TYPES, "threshold": THRESHOLD}
    if user["role"] == "student":
        ready = engine.is_ready(user["class_id"])
        data.update(classId=user["class_id"], ready=ready,
                    courses=engine.courses_for(user["class_id"]) if ready else [])
        if ready:
            cal = engine.calendar_for(user["class_id"])
            data["semester"] = {"start": cal["semesterStart"], "end": cal["semesterEnd"]}
    return data


# =========================================
# LEAVE APPLICATIONS
# =========================================


@app.post("/api/applications/preview")

def notify(conn, user_id: str, title: str, body: str, kind: str, application_id: str | None = None):
    conn.execute(
        """INSERT INTO notifications (id, user_id, title, body, kind, application_id, created_at)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (str(uuid.uuid4()), user_id, title, body, kind, application_id, now_iso()),
    )


def notify_faculty_new_app(conn, class_id: str, student_name: str, event_name: str, app_id: str,
                           title: str = "New application", kind: str = "new_application"):
    """Notify every faculty member who can see this class (its coordinators, and the HoD if any)."""
    faculty = conn.execute(
        """SELECT u.id FROM users u
           WHERE u.role = 'faculty' AND (
               EXISTS (SELECT 1 FROM faculty_classes f WHERE f.faculty_id = u.id AND f.class_id = ?)
               OR NOT EXISTS (SELECT 1 FROM faculty_classes f WHERE f.faculty_id = u.id))""",
        (class_id,),
    ).fetchall()
    for row in faculty:
        notify(conn, row["id"], title, f"{student_name}: {event_name}.", kind, app_id)


def preview(body: dict = Body(default_factory=dict), user=Depends(require_role("student"))):
    start_date, end_date = body.get("startDate"), body.get("endDate")
    result = validate_dates(user["class_id"], start_date, end_date)
    return {
        **result,
        "subjects": subject_summary(user["class_id"], result["lectures"]),
        "phases": engine.phases_for(user["class_id"]),
        "counted": body.get("kind") != "medical",
    }


@app.post("/api/applications", status_code=201)
async def submit_application(
    kind: str = Form("event"),
    category: str = Form(""),
    eventName: str = Form(""),
    startDate: str = Form(""),
    endDate: str = Form(""),
    description: str = Form(""),
    evidenceType: str = Form(""),
    approvalLetter: UploadFile | None = File(None),
    evidence: UploadFile | None = File(None),
    user=Depends(require_role("student")),
):
    kind, category = clean(kind), clean(category)
    event_name, description, evidence_type = clean(eventName), clean(description), clean(evidenceType)

    if kind not in ("event", "medical"):
        raise ApiError(400, "Choose event or medical.")
    if kind == "event":
        if category not in EVENT_CATEGORIES:
            raise ApiError(400, "Choose an event category.")
        if len(event_name) < 3:
            raise ApiError(400, "Enter the name of the event.")
        if evidence_type not in EVIDENCE_TYPES:
            raise ApiError(400, "Choose whether you have a certificate or HoD/Dean permission.")
    else:
        category, evidence_type = "Medical", "medical_certificate"
        event_name = event_name or "Medical leave"
    if len(description) > 500:
        raise ApiError(400, "Keep the description under 500 characters.")

    lectures = validate_dates(user["class_id"], startDate, endDate)["lectures"]
    if not lectures:
        raise ApiError(400, "No lectures fall on these dates, so there's nothing to add.")

    cert_pending = False
    files = []
    if kind == "event":
        if not approvalLetter or not approvalLetter.filename:
            raise ApiError(400, "Upload the pre-approval letter for the event.")
        files.append(("approval_letter", *await read_upload(approvalLetter, "pre-approval letter")))
        if evidence and evidence.filename:
            label = "certificate" if evidence_type == "certificate" else "HoD/Dean permission letter"
            files.append(("evidence", *await read_upload(evidence, label)))
        elif evidence_type == "authority_permission":
            raise ApiError(400, "Upload the HoD/Dean permission letter.")
        else:
            cert_pending = True   # certificate to be uploaded after the event
    else:
        files.append(("evidence", *await read_upload(evidence, "medical certificate")))

    with get_db() as conn:
        overlapping = conn.execute(
            """SELECT start_date, end_date FROM applications
               WHERE student_id = ? AND status != 'rejected' AND start_date <= ? AND end_date >= ?""",
            (user["id"], endDate, startDate),
        ).fetchone()
        if overlapping:
            raise ApiError(
                409,
                f"You already applied for {overlapping['start_date']} to {overlapping['end_date']}. "
                "Dates can't overlap.",
            )

        application_id = str(uuid.uuid4())
        submitted = now_iso()
        conn.execute(
            """INSERT INTO applications (id, student_id, class_id, kind, category, event_name,
                   description, start_date, end_date, lectures, evidence_type, certificate_pending,
                   submitted_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (application_id, user["id"], user["class_id"], kind, category, event_name, description,
             startDate, endDate, json.dumps(lectures), evidence_type, int(cert_pending), submitted),
        )
        notify_faculty_new_app(conn, user["class_id"], user["name"], event_name, application_id)

        written = []
        try:
            for doc_kind, content, upload in files:
                stored = f"{uuid.uuid4()}{ALLOWED_FILE_TYPES[upload.content_type]}"
                path = UPLOAD_DIR / stored
                path.write_bytes(content)
                written.append(path)
                conn.execute(
                    """INSERT INTO documents (id, application_id, kind, stored_name, original_name,
                           mime_type, size) VALUES (?, ?, ?, ?, ?, ?, ?)""",
                    (str(uuid.uuid4()), application_id, doc_kind, stored,
                     Path(upload.filename).name, upload.content_type, len(content)),
                )
        except Exception:
            for path in written:
                path.unlink(missing_ok=True)
            raise

        return {"application": application_json(get_application(conn, application_id), conn)}


@app.get("/api/applications")
def list_applications(classId: str | None = None, user=Depends(current_user)):
    with get_db() as conn:
        if user["role"] == "student":
            rows = conn.execute(
                "SELECT * FROM applications WHERE student_id = ? ORDER BY submitted_at DESC",
                (user["id"],),
            ).fetchall()
        else:
            allowed = visible_classes(conn, user)
            if classId and classId not in allowed:
                raise ApiError(403, "You don't coordinate this class.")
            wanted = [classId] if classId else allowed
            if not wanted:
                return {"applications": []}
            marks = ",".join("?" * len(wanted))
            rows = conn.execute(
                f"SELECT * FROM applications WHERE class_id IN ({marks}) ORDER BY submitted_at DESC",
                wanted,
            ).fetchall()
        return {"applications": [application_json(r, conn) for r in rows]}


@app.get("/api/applications/{application_id}/calculation")
def application_calculation(application_id: str, user=Depends(current_user)):
    """Sessions per subject for one application, split by CIA phase."""
    with get_db() as conn:
        row = get_application(conn, application_id)
        check_application_access(conn, user, row)
    lectures = json.loads(row["lectures"])
    return {
        "counted": row["kind"] == "event",
        "phases": engine.phases_for(row["class_id"]),
        "subjects": subject_summary(row["class_id"], lectures),
    }


@app.patch("/api/applications/{application_id}/read")
def mark_read(application_id: str, user=Depends(require_role("faculty"))):
    with get_db() as conn:
        row = get_application(conn, application_id)
        check_application_access(conn, user, row)
        conn.execute("UPDATE applications SET unread = 0 WHERE id = ?", (application_id,))
        return {"application": application_json(get_application(conn, application_id), conn)}


@app.patch("/api/applications/{application_id}/review")
def review(application_id: str, body: dict = Body(default_factory=dict),
           user=Depends(require_role("faculty"))):
    status, remark = clean(body.get("status")), clean(body.get("remark"))
    if status not in ("approved", "rejected"):
        raise ApiError(400, "Status must be approved or rejected.")
    if status == "rejected" and len(remark) < 3:
        raise ApiError(400, "Add a remark so the student knows why it was rejected.")

    with get_db() as conn:
        row = get_application(conn, application_id)
        check_application_access(conn, user, row)
        if row["status"] != "pending":
            raise ApiError(409, f"This application was already {row['status']}.")
        student_row = conn.execute("SELECT * FROM users WHERE id = ?", (row["student_id"],)).fetchone()
        conn.execute(
            """UPDATE applications SET status = ?, remark = ?, unread = 0, student_unread = 1,
                   reviewed_at = ?, reviewed_by = ? WHERE id = ?""",
            (status, remark, now_iso(), user["name"], application_id),
        )
        if student_row:
            verb = "approved" if status == "approved" else "rejected"
            notify(conn, student_row["id"], f"Application {verb}",
                   f'Your application for {row["event_name"]} has been {verb}.' +
                   (f" Note: {remark}" if remark else ""), "reviewed", application_id)
        return {"application": application_json(get_application(conn, application_id), conn)}


@app.get("/api/applications/{application_id}/documents/{document_id}")
def get_document(application_id: str, document_id: str, user=Depends(current_user)):
    with get_db() as conn:
        row = get_application(conn, application_id)
        check_application_access(conn, user, row)
        doc = conn.execute(
            "SELECT * FROM documents WHERE id = ? AND application_id = ?", (document_id, application_id)
        ).fetchone()
    if doc is None:
        raise ApiError(404, "Document not found.")
    path = UPLOAD_DIR / doc["stored_name"]
    if not path.exists():
        raise ApiError(404, "The document file is missing.")
    return FileResponse(path, media_type=doc["mime_type"])


# =========================================
# STUDENTS, ERP IMPORT AND REPORTS
# =========================================


@app.get("/api/students/{class_id}")
def list_students(class_id: str, user=Depends(require_role("faculty"))):
    with get_db() as conn:
        require_class(conn, user, class_id)
        students = class_students(conn, class_id)
    return {"students": [
        {"id": s["id"], **student_info(s), "activated": not s["must_change_password"]}
        for s in students
    ]}


@app.post("/api/reports/passwords/{class_id}")
def issue_passwords(class_id: str, user=Depends(require_role("faculty"))):
    """
    Creates new starting passwords for every student in the class who hasn't set their own
    password yet, and returns them as an Excel file. The passwords are saved and the file is
    built in the same request, so the file always matches the database.
    Students who already logged in and changed their password are not touched.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    with get_db() as conn:
        if not faculty_can_see(conn, user, class_id):
            raise ApiError(403, "You don't coordinate this class.")
        students = conn.execute(
            """SELECT * FROM users WHERE role = 'student' AND class_id = ?
               ORDER BY roll_no, name""",
            (class_id,),
        ).fetchall()
        if not students:
            raise ApiError(404, "No students in this class. Run import_data.py first.")

        issued, active = [], 0
        for s in students:
            if not s["must_change_password"]:
                active += 1  # already has their own password
                continue
            password = generate_password()
            conn.execute("UPDATE users SET password_hash = ? WHERE id = ?",
                         (hash_password(password), s["id"]))
            issued.append((s["roll_no"], s["prn"], s["name"], password))

    stamp = datetime.now()
    wb = Workbook()
    ws = wb.active
    ws.title = class_id
    ws.append([f"Starting passwords: {class_id}, created {stamp.strftime('%d-%m-%Y %H:%M')} by {user['name']}."])
    ws.append(["Only this file is valid. Any earlier file for these students no longer works. "
               "Give each student only their own row. Keep private."])
    ws.append(["Roll No.", "PRN (login ID)", "Name", "Starting password"])
    for row in issued:
        ws.append(list(row))
    if active:
        ws.append([])
        ws.append([f"{active} student(s) already set their own password and are not listed. "
                   "Use Reset in Class attendance if one of them forgets it."])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", bold=cell.row <= 3)
    for col, width in zip("ABCD", (9, 18, 36, 18)):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"starting_passwords_{class_id}_{stamp.strftime('%Y%m%d-%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"',
                 "X-Issued-Count": str(len(issued))},
    )


# ---------- ERP attendance import (class coordinator) ----------

ERP_HEADER_RE = re.compile(r"^\s*([A-Za-z0-9]+)\s+(attended|total)\s*$", re.IGNORECASE)
CLASS_NAMES = {"FY": "First Year", "SY-A": "Second Year", "SY-B": "DSY",
               "TY-A": "Third Year A", "TY-B": "Third Year B", "BTECH": "Final Year"}


def xlsx_response(wb, filename: str) -> StreamingResponse:
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def style_sheet(ws, widths, header_rows=3):
    """Flat report sheets: title rows, one header row, bordered data."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="B7C4D1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    last_col = max(ws.max_column, len(widths))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(1, header_rows):
        if ws.cell(row=r, column=1).value:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
    ws["A1"].font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", start_color="12355B")
    ws["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in range(2, header_rows):
        ws.cell(row=r, column=1).font = Font(name="Arial", italic=True, size=9, color="44546A")
        ws.cell(row=r, column=1).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 30
    for cell in ws[header_rows]:
        cell.font = Font(name="Arial", bold=True, size=10)
        cell.fill = PatternFill("solid", start_color="DCE6F1")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[header_rows].height = 32
    for row in ws.iter_rows(min_row=header_rows + 1):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="left" if cell.column in (3, 4) else "center")
            if cell.row <= ws.max_row and any(c.value not in (None, "") for c in row):
                cell.border = border
    ws.freeze_panes = ws.cell(row=header_rows + 1, column=4)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_rows}:{header_rows}"


def write_grid_sheet(ws, title: str, note: str, groups: list, students: list, values,
                     trailing: list | None = None, zero_as_dash: bool = False):
    """
    Grid sheet in the ERP-template layout:
      row 1  title (merged)          row 2  note (merged)
      row 3  subject names, each merged over its columns
      row 4  column headers (Roll No., PRN, Name, sub-columns...)
      row 5+ one row per student
    groups:   [(subject label, [column headers]), ...]
    values:   function(student_row, group_index, column_index) -> cell value
    trailing: [(header, function(row_number, group_column_letters) -> value or formula), ...]
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    trailing = trailing or []
    fixed = ["Roll No.", "PRN", "Name"]
    width = len(fixed) + sum(len(cols) for _, cols in groups) + len(trailing)
    thin = Side(style="thin", color="B7C4D1")
    medium = Side(style="medium", color="12355B")
    bands = ["EAF2FB", "FFFFFF"]
    band_heads = ["BDD7EE", "DDEBF7"]

    ws.cell(row=1, column=1, value=title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.cell(row=2, column=1, value=note)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A1"].font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", start_color="12355B")
    ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"].font = Font(name="Arial", italic=True, size=9, color="44546A")
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 30

    # Row 3: "Student" over the fixed columns, then each subject name over its columns
    ws.cell(row=3, column=1, value="Student")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=3)
    col = 4
    spans = []  # (first column, last column, band index) per group
    for gi, (label, cols) in enumerate(groups):
        ws.cell(row=3, column=col, value=label)
        if len(cols) > 1:
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + len(cols) - 1)
        spans.append((col, col + len(cols) - 1, gi % 2))
        col += len(cols)
    if trailing:
        ws.cell(row=3, column=col, value="Total")
        if len(trailing) > 1:
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + len(trailing) - 1)

    # Row 4: column headers
    headers = list(fixed)
    for _, cols in groups:
        headers += cols
    headers += [h for h, _ in trailing]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)

    # Rows 5+: students
    group_cols = [[get_column_letter(c) for c in range(a, b + 1)] for a, b, _ in spans]
    for r, s in enumerate(students, start=5):
        ws.cell(row=r, column=1, value=s["roll_no"])
        ws.cell(row=r, column=2, value=s["prn"])
        ws.cell(row=r, column=3, value=s["name"])
        c = 4
        for gi, (_, cols) in enumerate(groups):
            for ci in range(len(cols)):
                ws.cell(row=r, column=c, value=values(s, gi, ci))
                c += 1
        for _, fn in trailing:
            ws.cell(row=r, column=c, value=fn(r, group_cols))
            c += 1

    # Styling
    last_row = 4 + len(students)
    head_font = Font(name="Arial", bold=True, size=10)
    body_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=3, max_row=last_row, min_col=1, max_col=width):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if cell.row <= 4:
                cell.font = head_font
                cell.alignment = center
                cell.fill = PatternFill("solid", start_color="DCE6F1")
            else:
                cell.font = body_font
                cell.alignment = (Alignment(horizontal="left", vertical="center") if cell.column in (2, 3)
                                  else Alignment(horizontal="center", vertical="center"))
                if zero_as_dash and cell.column > 3:
                    cell.number_format = '0;-0;"-"'
    for a, b, band in spans:
        for row in ws.iter_rows(min_row=3, max_row=last_row, min_col=a, max_col=b):
            for cell in row:
                cell.fill = PatternFill("solid", start_color=(band_heads if cell.row <= 4 else bands)[band])
                left = medium if cell.column == a else thin
                cell.border = Border(left=left, right=thin, top=thin, bottom=thin)
    if trailing:
        first_trailing = width - len(trailing) + 1
        for row in ws.iter_rows(min_row=3, max_row=last_row, min_col=first_trailing, max_col=width):
            for cell in row:
                cell.fill = PatternFill("solid", start_color="FFF2CC" if cell.row > 4 else "FFE699")
                cell.font = Font(name="Arial", size=10, bold=True)
                left = medium if cell.column == first_trailing else thin
                cell.border = Border(left=left, right=thin, top=thin, bottom=thin)

    ws.row_dimensions[3].height = 42
    ws.row_dimensions[4].height = 30
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 32
    for i in range(4, width + 1):
        ws.column_dimensions[get_column_letter(i)].width = 11
    ws.freeze_panes = "D5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:4"


ERP_PROGRAM = "SET - B.Tech - CSE (Institute-Program)"


def write_erp_matrix(ws, title: str, info: str, groups: list, colnames: list, students: list,
                     cell_fn, grand_fn, trailing: list | None = None, zero_as_dash: bool = False):
    """
    Writes a sheet in the ERP "Division wise Subject wise Attendance %" layout:
      row 1 title, row 2 program|division|period, rows 3-4 empty,
      row 5 subject (merged), row 6 slot type (merged), row 7 PRN | Roll No | Name | column names,
      row 8+ one row per student, then a Grand Total block and any extra columns.
    groups:   [(subject label, [slot, ...]), ...]
    cell_fn:  (student, group index, slot, column name, row, {column name: letter}) -> value
    grand_fn: (student, column name, row, {column name: [letters]}, {column name: letter}) -> value
    trailing: [(header, (student, row, {column name: letter of the Grand Total block}) -> value), ...]
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    trailing = trailing or []
    thin, thick = Side(style="thin", color="9FB1C4"), Side(style="medium", color="12355B")
    col = 4
    layout = []  # (group index, slot, {name: letter}, first col)
    for gi, (_, slots) in enumerate(groups):
        for slot in slots:
            layout.append((gi, slot, {n: L(col + i) for i, n in enumerate(colnames)}, col))
            col += len(colnames)
    grand_first = col
    grand = {n: L(col + i) for i, n in enumerate(colnames)}
    col += len(colnames)
    trail_first = col
    last_col = col + len(trailing) - 1
    all_letters = {n: [lt[n] for _, _, lt, _ in layout] for n in colnames}

    title_end = min(last_col, 4 + 17)  # like ERP: title over the first subjects, visible without scrolling
    ws.cell(row=1, column=4, value=title)
    ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=title_end)
    ws.cell(row=2, column=4, value=info)
    ws.merge_cells(start_row=2, start_column=4, end_row=2, end_column=title_end)
    for r, size in ((1, 13), (2, 10)):
        c = ws.cell(row=r, column=4)
        c.font = Font(name="Arial", bold=True, size=size, color="12355B")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # Subject (row 5) and slot (row 6) headers
    for gi, (label, slots) in enumerate(groups):
        cols = [first for g, _, _, first in layout if g == gi]
        start, end = cols[0], cols[-1] + len(colnames) - 1
        ws.cell(row=5, column=start, value=label)
        if end > start:
            ws.merge_cells(start_row=5, start_column=start, end_row=5, end_column=end)
    for gi, slot, _, first in layout:
        ws.cell(row=6, column=first, value=slot)
        if len(colnames) > 1:
            ws.merge_cells(start_row=6, start_column=first, end_row=6, end_column=first + len(colnames) - 1)
    ws.cell(row=6, column=grand_first, value="Grand Total")
    if len(colnames) > 1:
        ws.merge_cells(start_row=6, start_column=grand_first, end_row=6, end_column=grand_first + len(colnames) - 1)
    if trailing:
        ws.cell(row=6, column=trail_first, value="Result")
        if len(trailing) > 1:
            ws.merge_cells(start_row=6, start_column=trail_first, end_row=6, end_column=last_col)

    # Column names (row 7)
    for i, h in enumerate(["PRN", "Roll No", "Name"], start=1):
        ws.cell(row=7, column=i, value=h)
    for _, _, letters, first in layout + [(None, None, grand, grand_first)]:
        for i, n in enumerate(colnames):
            ws.cell(row=7, column=first + i, value=n)
    for i, (h, _) in enumerate(trailing):
        ws.cell(row=7, column=trail_first + i, value=h)

    # Students
    for r, s in enumerate(students, start=8):
        ws.cell(row=r, column=1, value=s["prn"])
        ws.cell(row=r, column=2, value=s.get("roll_no"))
        ws.cell(row=r, column=3, value=s["name"])
        for gi, slot, letters, first in layout:
            for i, n in enumerate(colnames):
                ws.cell(row=r, column=first + i, value=cell_fn(s, gi, slot, n, r, letters))
        for i, n in enumerate(colnames):
            ws.cell(row=r, column=grand_first + i, value=grand_fn(s, n, r, all_letters, grand))
        for i, (_, fn) in enumerate(trailing):
            ws.cell(row=r, column=trail_first + i, value=fn(s, r, grand))

    # Styling
    last_row = 7 + len(students)
    bands = ["EAF2FB", "FFFFFF"]
    heads = ["BDD7EE", "DDEBF7"]
    for row in ws.iter_rows(min_row=5, max_row=last_row, min_col=1, max_col=last_col):
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            c.font = Font(name="Arial", size=10, bold=c.row <= 7)
            if c.row <= 7:
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c.fill = PatternFill("solid", start_color="DCE6F1")
            else:
                c.alignment = Alignment(horizontal="left" if c.column <= 3 else "center", vertical="center")
                if c.column > 3:
                    if zero_as_dash:
                        c.number_format = '0;-0;"-"'
                    elif ws.cell(row=7, column=c.column).value == "%":
                        c.number_format = "0.00"
    for gi, (_, slots) in enumerate(groups):
        cols = [first for g, _, _, first in layout if g == gi]
        start, end = cols[0], cols[-1] + len(colnames) - 1
        for row in ws.iter_rows(min_row=5, max_row=last_row, min_col=start, max_col=end):
            for c in row:
                c.fill = PatternFill("solid", start_color=(heads if c.row <= 7 else bands)[gi % 2])
                if c.column == start:
                    c.border = Border(left=thick, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=5, max_row=last_row, min_col=grand_first, max_col=last_col):
        for c in row:
            c.fill = PatternFill("solid", start_color="FFE699" if c.row <= 7 else "FFF2CC")
            c.font = Font(name="Arial", size=10, bold=True)
            if c.column in (grand_first, trail_first):
                c.border = Border(left=thick, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[5].height = 45
    ws.row_dimensions[7].height = 30
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 32
    for i in range(4, last_col + 1):
        ws.column_dimensions[L(i)].width = 10
    for i, (h, _) in enumerate(trailing):
        ws.column_dimensions[L(trail_first + i)].width = max(12, min(40, len(h) + 12))
    ws.freeze_panes = "D8"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = max(1, (last_col + 21) // 24)  # about 24 columns per printed page
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:7"
    ws.print_title_cols = "A:C"


def ref_sum(letters: list, row: int) -> str:
    return "=" + ("+".join(f"{x}{row}" for x in letters) if letters else "0")


@app.get("/api/erp/template/{class_id}")
def erp_template(class_id: str, user=Depends(require_role("faculty"))):
    """Excel sheet with every student and subject of the class, ready to fill from ERP."""
    from openpyxl import Workbook

    with get_db() as conn:
        require_class(conn, user, class_id)
        courses = engine.courses_for(class_id)
        students = class_students(conn, class_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "ERP attendance"
    write_grid_sheet(
        ws,
        title=f"ERP attendance: {class_id} ({len(students)} students)",
        note="Fill ATTENDED and TOTAL sessions for each subject exactly as in ERP. Leave both empty if a "
             "subject doesn't apply. Don't change PRNs or the column headers in row 4.",
        groups=[(f"{c['name']}\n{c['code']}", [f"{c['id']} Attended", f"{c['id']} Total"]) for c in courses],
        students=students,
        values=lambda s, gi, ci: None,
    )
    return xlsx_response(wb, f"ERP_attendance_template_{class_id}.xlsx")


def _read_sheet_rows(content: bytes, filename: str) -> list[list]:
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise ApiError(400, "Couldn't read the file. Upload the Excel template (.xlsx) or a .csv.")
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def _as_int(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError
    if isinstance(value, (int, float)):
        if float(value) != int(value):
            raise ValueError
        return int(value)
    text = str(value).strip()
    if not re.fullmatch(r"\d+(\.0+)?", text):
        raise ValueError
    return int(float(text))


def parse_template(rows: list[list], class_id: str) -> dict:
    """Our own ERP template (one 'Attended' and 'Total' column per subject), in the matrix shape."""
    header_index = next((i for i, r in enumerate(rows[:15])
                         if any(str(c).strip().upper() == "PRN" for c in r if c is not None)), None)
    if header_index is None:
        raise ApiError(400, "Couldn't find the header row with 'PRN'. Upload the ERP export as it is.")
    header = [str(c).strip() if c is not None else "" for c in rows[header_index]]
    prn_col = next(i for i, h in enumerate(header) if h.upper() == "PRN")
    courses = {c["id"].upper(): c for c in engine.courses_for(class_id)}
    columns = {}
    for i, h in enumerate(header):
        m = ERP_HEADER_RE.match(h)
        if not m:
            continue
        course = courses.get(m.group(1).upper())
        if course is None:
            raise ApiError(400, f"Column '{h}' isn't a subject of {class_id}.")
        columns.setdefault(course["id"], {})[m.group(2).lower()] = i
    columns = {c: v for c, v in columns.items() if "attended" in v and "total" in v}
    if not columns:
        raise ApiError(400, "This isn't the ERP 'Subject wise Attendance' export. Upload that file from ERP.")
    by_id = {c["id"]: c for c in engine.courses_for(class_id)}
    subjects = [{"key": cid, "code": by_id[cid]["code"], "name": by_id[cid]["name"],
                 "label": f"{by_id[cid]['code']} - {by_id[cid]['name']}", "slots": ["Total"]} for cid in columns]
    students, errors = [], []
    for n, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        prn = str(row[prn_col]).strip().upper() if prn_col < len(row) and row[prn_col] else ""
        if not prn:
            continue
        cells = {}
        for cid, cols in columns.items():
            try:
                attended = _as_int(row[cols["attended"]] if cols["attended"] < len(row) else None)
                total = _as_int(row[cols["total"]] if cols["total"] < len(row) else None)
            except ValueError:
                errors.append(f"Row {n} ({prn}), {cid}: use whole numbers.")
                continue
            if attended is None and total is None:
                continue
            if attended is None or total is None or attended > total:
                errors.append(f"Row {n} ({prn}), {cid}: check attended and total.")
                continue
            cells[(cid, "Total")] = [total, attended]
        students.append({"prn": prn, "cells": cells})
    return {"subjects": subjects, "students": students, "errors": errors,
            "periodFrom": None, "periodTo": None, "division": ""}


def dmy_to_iso(dmy: str | None) -> str | None:
    return datetime.strptime(dmy, "%d-%m-%Y").strftime("%Y-%m-%d") if dmy else None


@app.post("/api/erp/import/{class_id}", status_code=201)
async def erp_import(
    class_id: str,
    label: str = Form(""),
    asOf: str = Form(""),
    file: UploadFile | None = File(None),
    user=Depends(require_role("faculty")),
):
    """
    Imports the ERP export "Student Slot Type Wise Matrix" (Division wise Subject wise Attendance %)
    as it comes out of ERP (.xls, .xlsx or .csv). Our older template is still accepted.
    The name and as-on date are taken from the file when left empty.
    """
    if file is None or not file.filename:
        raise ApiError(400, "Choose the ERP attendance file.")
    content = await file.read(MAX_FILE_SIZE * 2 + 1)
    if len(content) > MAX_FILE_SIZE * 2:
        raise ApiError(400, "The file must be smaller than 10 MB.")
    try:
        rows = erp_matrix.read_rows(content, file.filename)
        parsed = erp_matrix.parse(rows) if erp_matrix.is_matrix(rows) else parse_template(rows, class_id)
    except erp_matrix.MatrixError as err:
        raise ApiError(400, str(err))

    if parsed["errors"]:
        errors = parsed["errors"]
        raise ApiError(400, f"{len(errors)} problem(s) in the file. Nothing was saved. "
                            + " ".join(errors[:6]) + (" ..." if len(errors) > 6 else ""))

    as_of = asOf or dmy_to_iso(parsed["periodTo"])
    if not as_of or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", as_of):
        raise ApiError(400, "Enter the date the ERP figures were taken on (the file doesn't say).")
    label = clean(label) or (f"ERP till {parsed['periodTo']}" if parsed["periodTo"] else "")
    if not 2 <= len(label) <= 40:
        raise ApiError(400, "Give this upload a name, e.g. Till CIA-1.")

    mapping = erp_matrix.map_to_courses(parsed["subjects"], engine.courses_for(class_id)) \
        if parsed["subjects"] and parsed["subjects"][0]["slots"] != ["Total"] \
        else {s["key"]: s["key"] for s in parsed["subjects"]}

    with get_db() as conn:
        require_class(conn, user, class_id)
        students = {s["prn"].upper(): s for s in class_students(conn, class_id)}
        known = [s for s in parsed["students"] if s["prn"] in students and s["cells"]]
        unknown = [s["prn"] for s in parsed["students"] if s["prn"] not in students]
        if not known:
            raise ApiError(400, f"None of the PRNs in this file belong to {class_id}. "
                                f"Is it the right division? The file says: {parsed['division'] or 'no division'}.")

        old = conn.execute("SELECT id FROM erp_imports WHERE class_id = ? AND label = ?",
                           (class_id, label)).fetchone()
        if old:
            conn.execute("DELETE FROM erp_imports WHERE id = ?", (old["id"],))
        import_id = str(uuid.uuid4())
        conn.execute(
            """INSERT INTO erp_imports (id, class_id, label, as_of, file_name, uploaded_by, uploaded_at,
                   division, period_from, period_to)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (import_id, class_id, label, as_of, Path(file.filename).name, user["name"], now_iso(),
             parsed["division"], dmy_to_iso(parsed["periodFrom"]), dmy_to_iso(parsed["periodTo"])),
        )
        conn.executemany(
            """INSERT INTO erp_subjects (import_id, ord, subject_key, code, name, label, slots, course_id)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            [(import_id, i, s["key"], s["code"], s["name"], s["label"], json.dumps(s["slots"]),
              mapping.get(s["key"])) for i, s in enumerate(parsed["subjects"])],
        )
        cells, records = [], []
        for s in known:
            sid = students[s["prn"]]["id"]
            per_subject: dict[str, list] = {}
            for (key, slot), (cond, pres) in s["cells"].items():
                cells.append((import_id, sid, key, slot, cond, pres))
                agg = per_subject.setdefault(key, [0, 0])
                agg[0] += cond
                agg[1] += pres
            records += [(import_id, sid, key, pres, cond) for key, (cond, pres) in per_subject.items()]
        conn.executemany("""INSERT INTO erp_cells (import_id, student_id, subject_key, slot, conducted, present)
                            VALUES (?, ?, ?, ?, ?, ?)""", cells)
        conn.executemany("INSERT INTO erp_records (import_id, student_id, course_id, attended, total) "
                         "VALUES (?, ?, ?, ?, ?)", records)

    seen = {s["prn"] for s in known}
    return {
        "importId": import_id,
        "replaced": bool(old),
        "label": label,
        "asOf": as_of,
        "division": parsed["division"],
        "periodFrom": dmy_to_iso(parsed["periodFrom"]),
        "periodTo": dmy_to_iso(parsed["periodTo"]),
        "students": len(seen),
        "subjects": len(parsed["subjects"]),
        "unmatchedSubjects": [s["label"] for s in parsed["subjects"] if s["key"] not in mapping],
        "unknownPrns": unknown[:20],
        "unknownCount": len(unknown),
        "missingStudents": len([p for p in students if p not in seen]),
    }


def list_imports(conn, class_id: str) -> list[dict]:
    rows = conn.execute(
        """SELECT i.*, COUNT(DISTINCT r.student_id) AS students
           FROM erp_imports i LEFT JOIN erp_records r ON r.import_id = i.id
           WHERE i.class_id = ? GROUP BY i.id ORDER BY i.as_of DESC, i.uploaded_at DESC""",
        (class_id,),
    ).fetchall()
    return [{"id": r["id"], "label": r["label"], "asOf": r["as_of"], "fileName": r["file_name"],
             "uploadedBy": r["uploaded_by"], "uploadedAt": r["uploaded_at"], "students": r["students"],
             "division": r["division"], "periodFrom": r["period_from"], "periodTo": r["period_to"]}
            for r in rows]


def load_erp(conn, import_id: str) -> tuple[list[dict], dict]:
    """
    Subjects of an ERP upload (in ERP order) and each student's numbers:
    ({key, code, name, label, slots, courseId}, {student_id: {(key, slot): (conducted, present)}}).
    Uploads made before the ERP matrix format are read from their per-subject totals.
    """
    subjects = [{"key": r["subject_key"], "code": r["code"], "name": r["name"], "label": r["label"],
                 "slots": json.loads(r["slots"]), "courseId": r["course_id"]}
                for r in conn.execute("SELECT * FROM erp_subjects WHERE import_id = ? ORDER BY ord", (import_id,))]
    cells: dict[str, dict] = {}
    if subjects:
        for r in conn.execute("SELECT * FROM erp_cells WHERE import_id = ?", (import_id,)):
            cells.setdefault(r["student_id"], {})[(r["subject_key"], r["slot"])] = (r["conducted"], r["present"])
        return subjects, cells
    rows = conn.execute("SELECT * FROM erp_records WHERE import_id = ?", (import_id,)).fetchall()
    imp = conn.execute("SELECT class_id FROM erp_imports WHERE id = ?", (import_id,)).fetchone()
    by_id = {c["id"]: c for c in engine.courses_for(imp["class_id"])} if imp else {}
    for key in dict.fromkeys(r["course_id"] for r in rows):
        c = by_id.get(key, {"code": "", "name": key})
        subjects.append({"key": key, "code": c["code"], "name": c["name"], "label": f"{c['code']} - {c['name']}",
                         "slots": ["Total"], "courseId": key})
    for r in rows:
        cells.setdefault(r["student_id"], {})[(r["course_id"], "Total")] = (r["total"], r["attended"])
    return subjects, cells


@app.get("/api/erp/imports/{class_id}")
def erp_imports(class_id: str, user=Depends(require_role("faculty"))):
    with get_db() as conn:
        require_class(conn, user, class_id)
        return {"imports": list_imports(conn, class_id)}


@app.delete("/api/erp/imports/{import_id}")
def delete_erp_import(import_id: str, user=Depends(require_role("faculty"))):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM erp_imports WHERE id = ?", (import_id,)).fetchone()
        if row is None:
            raise ApiError(404, "Upload not found.")
        require_class(conn, user, row["class_id"])
        conn.execute("DELETE FROM erp_imports WHERE id = ?", (import_id,))
    return {"deleted": True}


# ---------- Report 1: event attendance (approved applications, by CIA phase) ----------


def build_event_report(conn, class_id: str) -> dict:
    courses = engine.courses_for(class_id)
    phases = engine.phases_for(class_id)
    lectures_by_student = approved_event_lectures(conn, class_id)
    pending = pending_counts(conn, class_id)
    app_counts = {r["student_id"]: r["n"] for r in conn.execute(
        """SELECT student_id, COUNT(*) AS n FROM applications
           WHERE class_id = ? AND kind = 'event' AND status = 'approved' GROUP BY student_id""",
        (class_id,))}

    # Newest ERP upload (if any), so the report can show ERP attended next to granted sessions.
    imports = list_imports(conn, class_id)
    latest = imports[0] if imports else None
    erp: dict[str, dict] = {}
    if latest:
        erp_subjects, erp_cells = load_erp(conn, latest["id"])
        course_of = {x["key"]: x["courseId"] for x in erp_subjects}
        for student_id, cells in erp_cells.items():
            mine = erp.setdefault(student_id, {})
            for (key, _slot), (conducted, present) in cells.items():
                course = course_of.get(key)
                if not course:
                    continue
                agg = mine.setdefault(course, {"attended": 0, "total": 0})
                agg["attended"] += present
                agg["total"] += conducted

    students = []
    for s in class_students(conn, class_id):
        per_course = {c["id"]: {p["key"]: 0 for p in phases} for c in courses}
        by_phase = {p["key"]: 0 for p in phases}
        for lecture in lectures_by_student.get(s["id"], []):
            key = engine.phase_key(phases, lecture["date"])
            if key and lecture["courseId"] in per_course:
                per_course[lecture["courseId"]][key] += lecture["sessions"]
                by_phase[key] += lecture["sessions"]
        students.append({
            "student": {"id": s["id"], **student_info(s)},
            "perCourse": per_course,
            "byPhase": by_phase,
            "total": sum(by_phase.values()),
            "approvedApplications": app_counts.get(s["id"], 0),
            "pendingApplications": pending.get(s["id"], 0),
            "erp": erp.get(s["id"], {}),
        })
    return {"courses": courses, "phases": phases, "students": students, "erpImport": latest}


@app.get("/api/reports/event/{class_id}")
def event_report(class_id: str, user=Depends(require_role("faculty"))):
    with get_db() as conn:
        require_class(conn, user, class_id)
        return build_event_report(conn, class_id)


@app.get("/api/reports/event/{class_id}/xlsx")
def event_report_xlsx(class_id: str, user=Depends(require_role("faculty"))):
    """Event attendance in the ERP matrix layout: subject, Lab / Lecture, lectures granted per CIA phase."""
    from openpyxl import Workbook

    with get_db() as conn:
        require_class(conn, user, class_id)
        cls = conn.execute("SELECT * FROM classes WHERE id = ?", (class_id,)).fetchone()
        students = class_students(conn, class_id)
        lectures_by_student = approved_event_lectures(conn, class_id)
        imports = list_imports(conn, class_id)
        erp_subjects = load_erp(conn, imports[0]["id"])[0] if imports else []
        pending = pending_counts(conn, class_id)
        approved = {r["student_id"]: r["n"] for r in conn.execute(
            """SELECT student_id, COUNT(*) AS n FROM applications
               WHERE class_id = ? AND kind = 'event' AND status = 'approved' GROUP BY student_id""", (class_id,))}
        apps = conn.execute("SELECT * FROM applications WHERE class_id = ? ORDER BY submitted_at",
                            (class_id,)).fetchall()
        app_rows = [(a, conn.execute("SELECT * FROM users WHERE id = ?", (a["student_id"],)).fetchone(),
                     documents_for(conn, a["id"])) for a in apps]

    phases, courses = engine.phases_for(class_id), engine.courses_for(class_id)
    slots_of = engine.course_slots(class_id)
    erp_label = {x["courseId"]: x["label"] for x in reversed(erp_subjects) if x["courseId"]}
    groups = [(erp_label.get(c["id"], f"{c['code']} - {c['name']}"), slots_of.get(c["id"]) or ["Lecture"])
              for c in courses]

    # student id -> (course, slot, phase) -> sessions
    counts: dict[str, dict] = {}
    for sid, lectures in lectures_by_student.items():
        mine = counts.setdefault(sid, {})
        for lec in lectures:
            key = (lec["courseId"], engine.slot_of(lec.get("type", "TH")), engine.phase_key(phases, lec["date"]))
            mine[key] = mine.get(key, 0) + lec["sessions"]
    rows = [{"id": s["id"], "prn": s["prn"], "roll_no": s["roll_no"], "name": s["name"]} for s in students]
    cols = [p["label"] for p in phases] + ["Granted"]

    def cell(s, gi, slot, name, r, lt):
        if name == "Granted":
            return ref_sum([lt[p["label"]] for p in phases], r)
        phase = next(p["key"] for p in phases if p["label"] == name)
        return counts.get(s["id"], {}).get((courses[gi]["id"], slot, phase), 0)

    def grand(s, name, r, every, own):
        if name == "Granted":
            return ref_sum([own[p["label"]] for p in phases], r)
        return ref_sum(every[name], r)

    trailing = [("Approved applications", lambda s, r, g: approved.get(s["id"], 0)),
                ("Waiting for review", lambda s, r, g: pending.get(s["id"], 0))]

    cal = engine.calendar_for(class_id)
    info = (f"{ERP_PROGRAM}|{(imports[0]['division'] if imports and imports[0].get('division') else cls['label'])}"
            f"|{formatted(cal['semesterStart'])} to {formatted(cal['semesterEnd'])}")
    wb = Workbook()
    ws = wb.active
    ws.title = "Event attendance"
    write_erp_matrix(ws, "Division wise Subject wise Event Attendance (approved applications)", info,
                     groups, cols, rows, cell, grand, trailing, zero_as_dash=True)
    from openpyxl.styles import Font
    ws.cell(row=3, column=4, value=(
        "Lectures missed for approved events, from the time table (labs count as 1 session). "
        + ", ".join(f"{p['label']}: {formatted(p['start'])} to {formatted(p['end'])}" for p in phases)
        + f". Generated {datetime.now().strftime('%d-%m-%Y %H:%M')} by {user['name']}."))
    ws.cell(row=3, column=4).font = Font(name="Arial", italic=True, size=9, color="44546A")

    ws2 = wb.create_sheet("Applications and evidence")
    ws2.append([f"All applications: {class_id}"])
    ws2.append(["Medical applications are kept on record but never counted, as per the circular."])
    ws2.append(["Submitted", "Roll No.", "PRN", "Name", "Type", "Category", "Event", "From", "To",
                "Sessions", "Proof", "Documents", "Status", "Reviewed by", "Reviewed on", "Remark"])
    for a, s, docs in app_rows:
        lectures = json.loads(a["lectures"])
        ws2.append([
            formatted(a["submitted_at"][:10]), s["roll_no"] if s else "", s["prn"] if s else "",
            s["name"] if s else "", "Medical (not counted)" if a["kind"] == "medical" else "Event",
            a["category"], a["event_name"], formatted(a["start_date"]), formatted(a["end_date"]),
            sum(l["sessions"] for l in lectures),
            "Certificate pending" if a["certificate_pending"]
            else EVIDENCE_TYPES.get(a["evidence_type"], "Medical certificate"),
            "\n".join(f"{'Pre-approval letter' if d['kind'] == 'approval_letter' else 'Proof'}: {d['originalName']}"
                      for d in docs),
            a["status"].capitalize(), a["reviewed_by"] or "",
            formatted(a["reviewed_at"][:10]) if a["reviewed_at"] else "", a["remark"],
        ])
    if not app_rows:
        ws2.append(["No applications yet."])
    style_sheet(ws2, [12, 8, 16, 30, 18, 22, 28, 12, 12, 9, 30, 40, 11, 20, 12, 30])
    return xlsx_response(wb, f"Event_attendance_{class_id}_{datetime.now().strftime('%Y%m%d')}.xlsx")


def formatted(iso: str) -> str:
    """2026-08-14 -> 14-08-2026, the way dates are written in college documents."""
    return datetime.strptime(iso, "%Y-%m-%d").strftime("%d-%m-%Y")


# ---------- Report 2: final attendance (ERP + approved events) ----------


def allocate_events(subjects: list[dict], cells: dict, lectures: list[dict]) -> tuple[dict, int]:
    """
    Gives each missed event lecture to the matching ERP subject and slot (Lab / Lecture), capped at
    the sessions the student was absent in that slot (circular: credit only for lectures missed).
    Returns ({(key, slot): [event sessions, credit]}, event sessions with no matching ERP subject).
    """
    demand: dict[tuple, int] = {}
    for lec in lectures:
        k = (lec["courseId"], engine.slot_of(lec.get("type", "TH")))
        demand[k] = demand.get(k, 0) + lec["sessions"]
    result: dict[tuple, list] = {}
    unmatched = 0
    for (course, slot), sessions in demand.items():
        owned = [s for s in subjects if s["courseId"] == course]
        targets = [(s["key"], slot) for s in owned if (s["key"], slot) in cells]
        if not targets:  # e.g. the time table has a lab but ERP lists the subject only as a lecture
            targets = [(s["key"], sl) for s in owned for sl in s["slots"] if (s["key"], sl) in cells]
        if not targets:
            unmatched += sessions
            continue
        left = sessions
        for i, t in enumerate(targets):
            conducted, present = cells[t]
            slot_credit = result.setdefault(t, [0, 0])
            room = conducted - present - slot_credit[1]
            give = min(left, max(room, 0))
            slot_credit[1] += give
            slot_credit[0] += give if i < len(targets) - 1 else left  # remaining sessions stay visible on the last one
            left -= give
            if left == 0:
                break
    return result, unmatched


def build_final_report(conn, class_id: str, import_id: str | None) -> dict:
    imports = list_imports(conn, class_id)
    if not imports:
        raise ApiError(400, "Import the ERP attendance for this class first (ERP attendance page).")
    chosen = next((i for i in imports if i["id"] == import_id), None) if import_id else imports[0]
    if chosen is None:
        raise ApiError(404, "That ERP upload doesn't exist.")
    as_of = chosen["asOf"]

    subjects, erp_cells = load_erp(conn, chosen["id"])
    lectures_by_student = approved_event_lectures(conn, class_id)

    students, counts = [], {"detained": 0, "subject": 0, "clear": 0, "missing": 0}
    for s in class_students(conn, class_id):
        cells = erp_cells.get(s["id"])
        if not cells:
            counts["missing"] += 1
            students.append({"student": {"id": s["id"], **student_info(s)}, "subjects": [],
                             "erpPercent": None, "finalPercent": None, "credit": 0, "unmatched": 0,
                             "below": [], "erpBelow": [], "status": "No ERP data", "statusKey": "missing"})
            continue
        # Only event lectures up to the ERP date: later sessions aren't in the ERP numbers yet.
        lectures = [l for l in lectures_by_student.get(s["id"], []) if l["date"] <= as_of]
        granted, unmatched = allocate_events(subjects, cells, lectures)

        rows = []
        for subj in subjects:
            slots = {}
            for slot in subj["slots"]:
                if (subj["key"], slot) not in cells:
                    continue
                conducted, present = cells[(subj["key"], slot)]
                events, credit = granted.get((subj["key"], slot), [0, 0])
                slots[slot] = {"conducted": conducted, "present": present, "eventSessions": events, "granted": credit}
            if not slots:
                continue
            total = sum(v["conducted"] for v in slots.values())
            attended = sum(v["present"] for v in slots.values())
            credit = sum(v["granted"] for v in slots.values())
            rows.append({
                "courseId": subj["key"], "subjectCode": subj["code"], "subjectName": subj["name"],
                "attended": attended, "total": total, "erpPercent": pct(attended, total) or 0.0,
                "eventSessions": sum(v["eventSessions"] for v in slots.values()), "credit": credit,
                "finalPercent": pct(attended + credit, total) or 0.0, "slots": slots,
            })

        sum_att = sum(r["attended"] for r in rows)
        sum_total = sum(r["total"] for r in rows)
        sum_credit = sum(r["credit"] for r in rows)
        final = pct(sum_att + sum_credit, sum_total) or 0.0
        erp_pct = pct(sum_att, sum_total) or 0.0
        below = [r["subjectName"] for r in rows if r["total"] and r["finalPercent"] < THRESHOLD]
        erp_below = [r["subjectName"] for r in rows if r["total"] and r["erpPercent"] < THRESHOLD]
        if final < THRESHOLD:
            key, status = "detained", "Detained"
        elif below:
            key, status = "subject", f"Detained in {len(below)} subject{'s' if len(below) > 1 else ''}"
        else:
            key, status = "clear", "Not detained"
        counts[key] += 1
        students.append({
            "student": {"id": s["id"], **student_info(s)}, "subjects": rows,
            "erpPercent": erp_pct, "finalPercent": final, "credit": sum_credit, "unmatched": unmatched,
            "below": below, "erpBelow": erp_below, "status": status, "statusKey": key,
            "savedByEvents": (erp_pct < THRESHOLD or bool(erp_below)) and key == "clear",
        })

    return {"import": chosen, "imports": imports, "courses": engine.courses_for(class_id),
            "erpSubjects": subjects, "threshold": THRESHOLD, "counts": counts, "students": students}


@app.get("/api/reports/final/{class_id}")
def final_report(class_id: str, importId: str | None = None, user=Depends(require_role("faculty"))):
    with get_db() as conn:
        require_class(conn, user, class_id)
        return build_final_report(conn, class_id, importId)


@app.get("/api/reports/final/{class_id}/xlsx")
def final_report_xlsx(class_id: str, importId: str | None = None, user=Depends(require_role("faculty"))):
    """Final attendance in the ERP matrix layout: every subject and slot, with the granted event lectures."""
    from openpyxl import Workbook

    with get_db() as conn:
        require_class(conn, user, class_id)
        report = build_final_report(conn, class_id, importId)
        cls = conn.execute("SELECT * FROM classes WHERE id = ?", (class_id,)).fetchone()
    imp, subjects = report["import"], report["erpSubjects"]
    by_student = {r["student"]["id"]: r for r in report["students"]}
    students = [{"prn": r["student"]["prn"], "roll_no": r["student"].get("rollNo"), "name": r["student"]["name"],
                 "id": r["student"]["id"]} for r in report["students"]]
    period_from = formatted(imp["periodFrom"]) if imp.get("periodFrom") else "semester start"
    info = f"{ERP_PROGRAM}|{imp.get('division') or cls['label']}|{period_from} to {formatted(imp['asOf'])}"
    cols = ["Conducted", "Present", "Granted", "Total Present", "Absent", "%"]

    def slot_data(s, gi, slot):
        row = by_student[s["id"]]
        subj = next((x for x in row["subjects"] if x["courseId"] == subjects[gi]["key"]), None)
        return subj["slots"].get(slot) if subj else None

    def cell(s, gi, slot, name, r, lt):
        d = slot_data(s, gi, slot)
        if d is None:
            return None
        return {
            "Conducted": d["conducted"],
            "Present": d["present"],
            "Granted": d["granted"],
            "Total Present": f"={lt['Present']}{r}+{lt['Granted']}{r}",
            "Absent": f"={lt['Conducted']}{r}-{lt['Total Present']}{r}",
            "%": f'=IF({lt["Conducted"]}{r}>0,ROUND({lt["Total Present"]}{r}/{lt["Conducted"]}{r}*100,2),"")',
        }[name]

    def grand(s, name, r, every, own):
        if by_student[s["id"]]["statusKey"] == "missing":
            return None
        if name in ("Conducted", "Present", "Granted"):
            return ref_sum(every[name], r)
        return {
            "Total Present": f"={own['Present']}{r}+{own['Granted']}{r}",
            "Absent": f"={own['Conducted']}{r}-{own['Total Present']}{r}",
            "%": f'=IF({own["Conducted"]}{r}>0,ROUND({own["Total Present"]}{r}/{own["Conducted"]}{r}*100,2),"")',
        }[name]

    trailing = [
        ("ERP %", lambda s, r, g: None if by_student[s["id"]]["statusKey"] == "missing"
            else f'=IF({g["Conducted"]}{r}>0,ROUND({g["Present"]}{r}/{g["Conducted"]}{r}*100,2),"")'),
        ("Subjects below 75%", lambda s, r, g: ", ".join(by_student[s["id"]]["below"])),
        ("Status", lambda s, r, g: by_student[s["id"]]["status"]),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Final attendance"
    write_erp_matrix(
        ws, "Division wise Subject wise Attendance % (ERP + granted event attendance)", info,
        [(x["label"], x["slots"]) for x in subjects], cols, students, cell, grand, trailing,
    )
    ws.cell(row=3, column=4, value=(
        f"ERP upload '{imp['label']}' as on {formatted(imp['asOf'])} plus lectures missed for approved events "
        f"up to that date. Granted is capped at the lectures actually missed. Minimum required: {THRESHOLD:g}%. "
        f"Generated {datetime.now().strftime('%d-%m-%Y %H:%M')} by {user['name']}."))
    from openpyxl.styles import Font
    ws.cell(row=3, column=4).font = Font(name="Arial", italic=True, size=9, color="44546A")

    ws2 = wb.create_sheet("Summary")
    ws2.append([f"Final attendance summary: {class_id} ({imp.get('division') or cls['label']})"])
    ws2.append([f"ERP as on {formatted(imp['asOf'])} + approved event attendance. Minimum {THRESHOLD:g}%."])
    ws2.append(["Roll No.", "PRN", "Name", "Conducted", "ERP present", "Granted", "Total present",
                "ERP %", "Final %", "Subjects below 75%", "Status"])
    for r in report["students"]:
        n = ws2.max_row + 1
        if r["statusKey"] == "missing":
            ws2.append([r["student"].get("rollNo"), r["student"]["prn"], r["student"]["name"],
                        "", "", "", "", "", "", "", "No ERP data"])
            continue
        total = sum(x["total"] for x in r["subjects"])
        present = sum(x["attended"] for x in r["subjects"])
        ws2.append([r["student"].get("rollNo"), r["student"]["prn"], r["student"]["name"], total, present,
                    r["credit"], f"=E{n}+F{n}", f'=IF(D{n}>0,ROUND(E{n}/D{n}*100,2),"")',
                    f'=IF(D{n}>0,ROUND(G{n}/D{n}*100,2),"")', ", ".join(r["below"]), r["status"]])
        ws2[f"H{n}"].number_format = ws2[f"I{n}"].number_format = "0.00"
    style_sheet(ws2, [8, 16, 32, 11, 11, 10, 12, 9, 9, 40, 22])
    return xlsx_response(wb, f"Final_attendance_{class_id}_{imp['asOf']}.xlsx")


def detention_docx(entries: list[tuple[str, dict]], prepared_by: str, skipped: list[str]) -> io.BytesIO:
    """
    The detention list in the department's Word format (letterhead from seed/detention_template.docx).
    entries: [(class_id, final report), ...] in the order they should appear.
    """
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    template = Path(__file__).parent / "seed" / "detention_template.docx"
    doc = docx.Document(str(template)) if template.exists() else docx.Document()
    for p in (doc.sections[0].header.paragraphs if template.exists() else []):
        if "Academic Year" in p.text and p.runs:
            p.runs[0].text = "                                    ODD SEMESTER, Academic Year 2026-27"
            for extra in p.runs[1:]:
                extra.text = ""

    def para(text="", bold=False, size=12, align=None, space_after=4):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold, run.font.size = bold, Pt(size)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        return p

    day = datetime.now()
    suffix = "th" if 11 <= day.day <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(day.day % 10, "th")
    para(f"Date: {day.day}{suffix} {day.strftime('%B %Y')}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    para("Detention List", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    para("This is to inform that, as per the university attendance criteria, the following students have not "
         "maintained the minimum required attendance and are therefore detained.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para("As per university rules, students who fail to meet the minimum attendance requirement are not eligible "
         "to appear for the End Semester Examination (ESE) scheduled as per the academic calendar.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para("Note : These students will not be permitted to appear for the above examinations under any "
         "circumstances. Students are advised to take this matter seriously and comply with university "
         "regulations.", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=10)

    table = doc.add_table(rows=1, cols=6)
    borders = OxmlElement("w:tblBorders")  # the letterhead template has no "Table Grid" style
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "000000")):
            el.set(qn(key), value)
        borders.append(el)
    table._tbl.tblPr.append(borders)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ["Sr. No", "Class", "Roll No.", "Name of the Student",
                                                "Attendance", "Remark"]):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True

    sr = 0
    for class_id, report in entries:
        for r in report["students"]:
            if r["statusKey"] not in ("detained", "subject"):
                continue
            sr += 1
            if r["statusKey"] == "detained":
                remark = "Permanently Detained" if (r["finalPercent"] or 0) == 0 else "Detained"
            else:
                n = len(r["below"])
                remark = f"{n} Subject{'s' if n > 1 else ''} detained ({', '.join(r['below'])})"
            values = [f"{sr}.", CLASS_NAMES.get(class_id, class_id), str(r["student"]["rollNo"] or ""),
                      r["student"]["name"], f"{r['finalPercent']:g}%", remark]
            for cell, text in zip(table.add_row().cells, values):
                cell.text = text
    if sr == 0:
        cells = table.add_row().cells
        cells[0].merge(cells[5]).text = "No students are detained."

    # Repeat the header row on every page and keep each row on one page.
    for i, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))

    table.autofit = False
    widths = [Inches(w) for w in (0.55, 1.0, 0.6, 2.0, 0.9, 1.5)]
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)

    para()
    basis = "; ".join(f"{cid}: ERP as on {formatted(rep['import']['asOf'])}" for cid, rep in entries)
    para(f"Attendance = ERP attendance plus lectures missed for approved events ({basis}).", size=9)
    if skipped:
        para(f"Not included, no ERP attendance uploaded yet: {', '.join(skipped)}.", size=9)
    para()
    para("Dr. Mahendra Gawali", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    para("HoD CSE", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(f"Prepared by: {prepared_by}", size=10)
    para("Copy to", bold=True)
    for line in ("Registrar Office", "SET DEAN Office", "Class Coordinator File"):
        para(line, space_after=0)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def docx_response(buffer: io.BytesIO, filename: str) -> StreamingResponse:
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/reports/final/{class_id}/detention-list")
def detention_list_docx(class_id: str, importId: str | None = None, user=Depends(require_role("faculty"))):
    """Detention list for one class."""
    with get_db() as conn:
        require_class(conn, user, class_id)
        report = build_final_report(conn, class_id, importId)
    buffer = detention_docx([(class_id, report)], f"{user['name']}, Class Coordinator ({class_id})", [])
    return docx_response(buffer, f"Detention_List_{class_id}_{datetime.now().strftime('%Y%m%d')}.docx")


@app.get("/api/reports/detention-list")
def department_detention_list(user=Depends(require_role("faculty"))):
    """
    One detention list for every class this faculty member can see (HoD: the whole department),
    like the department's own list. Each class uses its newest ERP upload.
    """
    entries, skipped = [], []
    with get_db() as conn:
        for class_id in visible_classes(conn, user):
            if not engine.is_ready(class_id) or not list_imports(conn, class_id):
                skipped.append(class_id)
                continue
            entries.append((class_id, build_final_report(conn, class_id, None)))
    if not entries:
        raise ApiError(400, "No class has ERP attendance uploaded yet, so there is nothing to list.")
    who = f"{user['name']}, {user['designation'] or 'Faculty'}"
    buffer = detention_docx(entries, who, skipped)
    return docx_response(buffer, f"Detention_List_CSE_{datetime.now().strftime('%Y%m%d')}.docx")



# =========================================
# NOTIFICATIONS AND REMINDER
# =========================================

@app.get("/api/notifications")
def get_notifications(user=Depends(current_user)):
    with get_db() as conn:
        rows = conn.execute(
            """SELECT * FROM notifications WHERE user_id = ?
               ORDER BY created_at DESC LIMIT 30""",
            (user["id"],),
        ).fetchall()
        unread = conn.execute(
            "SELECT COUNT(*) FROM notifications WHERE user_id = ? AND read_at IS NULL",
            (user["id"],),
        ).fetchone()[0]
    return {
        "unread": unread,
        "notifications": [
            {"id": r["id"], "title": r["title"], "body": r["body"], "kind": r["kind"],
             "applicationId": r["application_id"], "createdAt": r["created_at"],
             "readAt": r["read_at"]}
            for r in rows
        ],
    }


@app.patch("/api/notifications/{notification_id}/read")
def mark_notification_read(notification_id: str, user=Depends(current_user)):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM notifications WHERE id = ?", (notification_id,)).fetchone()
        if not row or row["user_id"] != user["id"]:
            raise ApiError(404, "Notification not found.")
        conn.execute("UPDATE notifications SET read_at = ? WHERE id = ?", (now_iso(), notification_id))
    return {"ok": True}


@app.patch("/api/notifications/read-all")
def mark_all_read(user=Depends(current_user)):
    with get_db() as conn:
        conn.execute("UPDATE notifications SET read_at = ? WHERE user_id = ? AND read_at IS NULL",
                     (now_iso(), user["id"]))
    return {"ok": True}


@app.patch("/api/applications/{application_id}/student-read")
def student_mark_read(application_id: str, user=Depends(require_role("student"))):
    with get_db() as conn:
        row = get_application(conn, application_id)
        if row["student_id"] != user["id"]:
            raise ApiError(403, "Not your application.")
        conn.execute("UPDATE applications SET student_unread = 0 WHERE id = ?", (application_id,))
    return {"ok": True}


@app.post("/api/applications/{application_id}/reminder")
def send_reminder(application_id: str, user=Depends(require_role("student"))):
    with get_db() as conn:
        row = get_application(conn, application_id)
        if row["student_id"] != user["id"]:
            raise ApiError(403, "Not your application.")
        if row["status"] != "pending":
            raise ApiError(409, f"This application was already {row['status']}.")
        submitted = row["submitted_at"]
        age = datetime.now(timezone.utc) - datetime.fromisoformat(submitted.replace("Z", "+00:00"))
        if age.total_seconds() < 86400:
            hours_left = int((86400 - age.total_seconds()) / 3600) + 1
            raise ApiError(429, f"Reminders can be sent 24 hours after submitting. "
                                f"Try again in about {hours_left} hour{'s' if hours_left != 1 else ''}.")
        last = row["reminder_sent_at"]
        if last:
            since = datetime.now(timezone.utc) - datetime.fromisoformat(last.replace("Z", "+00:00"))
            if since.total_seconds() < 86400:
                raise ApiError(429, "You already sent a reminder in the last 24 hours.")
        conn.execute("UPDATE applications SET reminder_sent_at = ? WHERE id = ?",
                     (now_iso(), application_id))
        notify_faculty_new_app(conn, row["class_id"], user["name"],
                               f"waiting for review since {row['submitted_at'][:10]} ({row['event_name']})",
                               application_id, title="Reminder from student", kind="reminder")
    return {"ok": True, "eventName": row["event_name"]}


@app.patch("/api/applications/{application_id}/upload-certificate")
async def upload_certificate(
    application_id: str,
    file: UploadFile | None = File(None),
    user=Depends(require_role("student")),
):
    if file is None or not file.filename:
        raise ApiError(400, "Choose the certificate file.")
    content, upload = await read_upload(file, "participation certificate")
    with get_db() as conn:
        row = get_application(conn, application_id)
        if row["student_id"] != user["id"]:
            raise ApiError(403, "Not your application.")
        if row["status"] == "rejected":
            raise ApiError(409, "This application was rejected and can't be updated.")
        if not row["certificate_pending"]:
            raise ApiError(409, "A certificate was already uploaded for this application.")
        stored = f"{uuid.uuid4()}{ALLOWED_FILE_TYPES[upload.content_type]}"
        (UPLOAD_DIR / stored).write_bytes(content)
        conn.execute(
            """INSERT INTO documents (id, application_id, kind, stored_name, original_name,
                   mime_type, size) VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (str(uuid.uuid4()), application_id, "evidence", stored,
             Path(upload.filename).name, upload.content_type, len(content)),
        )
        conn.execute("UPDATE applications SET certificate_pending = 0, unread = 1 WHERE id = ?",
                     (application_id,))
        notify_faculty_new_app(conn, row["class_id"], user["name"],
                               f"certificate uploaded for {row['event_name']}", application_id,
                               title="Certificate uploaded")
        return {"application": application_json(get_application(conn, application_id), conn)}


if __name__ == "__main__":
    print(f"API running on http://localhost:{PORT}  (docs: http://localhost:{PORT}/docs)")
    uvicorn.run("main:app", host="127.0.0.1", port=PORT, reload=False)
