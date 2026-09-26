"""
Reads the ERP export "Student Slot Type Wise Matrix" (Division wise Subject wise Attendance %).

Layout of the ERP file:
  row 1   title                                  "Division wise Subject wise Attendance %"
  row 2   program | division | period            "SET - B.Tech - CSE ...|SET - FY CSE - A|13-07-2026 to 15-09-2026"
  row 5   subject, merged over its columns       "26UCSES102 - Introduction of Data Structure using C++"
  row 6   slot type, merged over its columns     "Lab" / "Lecture" / "Tutorial" / "Grand Total"
  row 7   PRN | Roll No | Name | Conducted | Present | Absent | % | ...
  row 8+  one row per student; empty cells = subject doesn't apply to that student

The same subject can appear more than once (with and without the code, or as an abbreviation);
those blocks are merged into one subject. Works with .xls (ERP default), .xlsx and .csv.
"""

import csv
import io
import re

CODE_RE = re.compile(r"^\s*(\d{2}[A-Z]{3,}[A-Z0-9]*\d[A-Z0-9]*|\d{2}[A-Z]{3,}X+)\s*-\s*(.+?)\s*$")
PAREN_RE = re.compile(r"\(\s*([A-Za-z]{2,10})\s*\)\s*$")
SMALL_WORDS = {"of", "and", "the", "using", "in", "for", "to", "with", "a", "an"}
SLOT_NAMES = {"lab": "Lab", "lecture": "Lecture", "tutorial": "Tutorial", "practical": "Lab", "theory": "Lecture"}


class MatrixError(ValueError):
    pass


def read_rows(content: bytes, filename: str) -> list[list]:
    """All cells of the first sheet as a list of rows. Merged cells: only the first cell has the value."""
    name = filename.lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return [list(r) for r in csv.reader(io.StringIO(text))]
    if name.endswith(".xls"):
        import xlrd
        try:
            book = xlrd.open_workbook(file_contents=content, logfile=io.StringIO())
        except Exception:
            raise MatrixError("Couldn't read the .xls file. Export it again from ERP, or save it as .xlsx.")
        sheet = book.sheet_by_index(0)
        return [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise MatrixError("Couldn't read the file. Upload the ERP export (.xls or .xlsx) or a .csv.")
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def text(v) -> str:
    return "" if v is None else str(v).strip()


def number(v):
    """Whole number or None for an empty cell."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        raise ValueError(v)
    if f < 0 or f != int(f):
        raise ValueError(v)
    return int(f)


def norm(s: str) -> str:
    s = re.sub(r"[^a-z0-9]+", " ", s.lower())
    return " ".join(s.split())


def initials(name: str) -> str:
    words = [w for w in norm(name).split() if w not in SMALL_WORDS and not w.isdigit()]
    return "".join(w[0] for w in words).upper()


def is_matrix(rows: list[list]) -> bool:
    """True if this looks like the ERP matrix (has 'Conducted' headers)."""
    return any(text(c).lower() == "conducted" for row in rows[:15] for c in row)


def parse(rows: list[list]) -> dict:
    header_i = next((i for i, r in enumerate(rows[:15])
                     if any(text(c).upper() == "PRN" for c in r) and any(text(c).lower() == "conducted" for c in r)),
                    None)
    if header_i is None or header_i < 2:
        raise MatrixError("This doesn't look like the ERP 'Subject wise Attendance' export: "
                          "no header row with PRN and Conducted.")
    header = [text(c) for c in rows[header_i]]
    subj_row = [text(c) for c in rows[header_i - 2]]
    slot_row = [text(c) for c in rows[header_i - 1]]
    width = len(header)
    subj_row += [""] * (width - len(subj_row))
    slot_row += [""] * (width - len(slot_row))

    prn_col = next(i for i, h in enumerate(header) if h.upper() == "PRN")
    name_col = next((i for i, h in enumerate(header) if h.lower() == "name"), None)

    # Walk the columns, carrying merged subject / slot labels forward.
    blocks = []  # {label, slot, cols: {conducted, present}}
    grand = {}
    subject, slot = "", ""
    for c in range(width):
        if subj_row[c]:
            subject, slot = subj_row[c], ""
        if slot_row[c]:
            slot = slot_row[c]
        h = header[c].lower()
        if h not in ("conducted", "present"):
            continue
        if slot.lower() == "grand total":
            grand[h] = c
            continue
        if not subject:
            continue
        slot_name = SLOT_NAMES.get(slot.lower(), slot or "Lecture")
        if not blocks or blocks[-1]["label"] != subject or blocks[-1]["slot"] != slot_name or h in blocks[-1]["cols"]:
            blocks.append({"label": subject, "slot": slot_name, "cols": {}})
        blocks[-1]["cols"][h] = c
    blocks = [b for b in blocks if "conducted" in b["cols"] and "present" in b["cols"]]
    if not blocks:
        raise MatrixError("No subject columns (Conducted / Present) found in the ERP file.")

    # Merge blocks of the same subject.
    subjects: dict[str, dict] = {}
    aliases: dict[str, str] = {}

    def add_subject(key, code, name, label):
        if key not in subjects:
            subjects[key] = {"key": key, "code": code, "name": name, "label": label, "slots": []}
        return subjects[key]

    for b in blocks:  # coded subjects first, so name-only copies can find them
        m = CODE_RE.match(b["label"])
        if not m:
            continue
        code, rest = m.group(1), m.group(2)
        abbrev = PAREN_RE.search(rest)
        name = PAREN_RE.sub("", rest).strip()
        key = code[2:]  # drop the batch year: 26UCSES102 -> UCSES102
        add_subject(key, code, name, f"{code} - {name}")
        b["key"] = key
        for alias in (norm(name), initials(name), abbrev.group(1).upper() if abbrev else None):
            if alias:
                aliases.setdefault(alias, key)
    for b in blocks:
        if "key" in b:
            continue
        label = b["label"]
        key = aliases.get(norm(label)) or aliases.get(label.upper().replace(" ", "")) or aliases.get(initials(label))
        if key is None:
            key = "N:" + norm(label)
            add_subject(key, "", label, label)
        b["key"] = key
    for b in blocks:
        if b["slot"] not in subjects[b["key"]]["slots"]:
            subjects[b["key"]]["slots"].append(b["slot"])

    # Students
    students, errors = [], []
    for n, row in enumerate(rows[header_i + 1:], start=header_i + 2):
        row = list(row) + [None] * (width - len(row))
        prn = text(row[prn_col]).upper()
        if not prn:
            continue
        cells: dict[tuple, list] = {}
        for b in blocks:
            try:
                cond = number(row[b["cols"]["conducted"]])
                pres = number(row[b["cols"]["present"]])
            except ValueError:
                errors.append(f"Row {n} ({prn}), {b['label']} {b['slot']}: not a whole number.")
                continue
            if cond is None and pres is None:
                continue
            cond, pres = cond or 0, pres or 0
            if pres > cond:
                errors.append(f"Row {n} ({prn}), {b['label']} {b['slot']}: present {pres} > conducted {cond}.")
                continue
            cell = cells.setdefault((b["key"], b["slot"]), [0, 0])
            cell[0] += cond
            cell[1] += pres
        grand_ok = None
        if "conducted" in grand and "present" in grand:
            try:
                gc, gp = number(row[grand["conducted"]]), number(row[grand["present"]])
                if gc is not None:
                    grand_ok = (gc, gp) == (sum(v[0] for v in cells.values()), sum(v[1] for v in cells.values()))
            except ValueError:
                pass
        students.append({"prn": prn, "name": text(row[name_col]) if name_col is not None else "",
                         "cells": cells, "grandMatches": grand_ok})

    # Title / division / period (row 2 in the ERP file)
    info = " ".join(text(c) for r in rows[:header_i - 2] for c in r if text(c))
    period = re.search(r"(\d{2}-\d{2}-\d{4})\s*to\s*(\d{2}-\d{2}-\d{4})", info)
    parts = [p.strip() for p in info.split("|")]
    division = parts[1] if len(parts) >= 3 else ""
    return {
        "subjects": list(subjects.values()),
        "students": students,
        "errors": errors,
        "periodFrom": period.group(1) if period else None,
        "periodTo": period.group(2) if period else None,
        "division": division,
        "title": text(rows[0][next((i for i, c in enumerate(rows[0]) if text(c)), 0)]) if rows and rows[0] else "",
    }


def map_to_courses(subjects: list[dict], courses: list[dict]) -> dict[str, str]:
    """ERP subject key -> our time-table course id (only confident, unique matches)."""
    result = {}

    def unique(pred):
        hits = [c["id"] for c in courses if pred(c)]
        return hits[0] if len(hits) == 1 else None

    for s in subjects:
        code, name = s["code"], s["name"]
        found = None
        if code:
            found = (unique(lambda c: c.get("code", "")[2:] == code[2:])
                     or unique(lambda c: len(c.get("code", "")) >= 5 and c["code"][-5:] == code[-5:]))
        found = (found
                 or unique(lambda c: norm(c["name"]) == norm(name))
                 or unique(lambda c: c["id"].upper() in (initials(name), name.upper().replace(" ", "")))
                 or unique(lambda c: initials(c["name"]) == initials(name) and len(initials(name)) >= 3))
        if found:
            result[s["key"]] = found
    return result
